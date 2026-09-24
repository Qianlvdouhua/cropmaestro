"""Loss-aware tabular input and deterministic field profiling."""

from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, localcontext
import csv
import hashlib
import io
from pathlib import Path
import posixpath
import re
import xml.etree.ElementTree as ET
import zipfile


RESERVED = {'query_task', 'create_at', 'created_at', 'query_hash', 'session_id'}
NUMBER = re.compile(r'^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$', re.ASCII)
IDENTIFIER = re.compile(r'编号|编码|代码|标识|(?:^|_)(?:id|code|uuid|doi)(?:_|$)|accession|accenumb|instcode', re.I)
CATEGORY = re.compile(r'等级|颜色|色泽|形状|株型|粒型|穗形|类型|抗性|color|colour|shape|attitude|presence|resistance|pubescence|branching|distribution|arrangement|exsertion|shattering|评分代码', re.I)
COUNT = re.compile(r'数(?:量)?$|(?:^|_)(?:number|count)(?:_|$)', re.I)
CROP_NAMES = {'rice': 'rice', '水稻': 'rice', '稻': 'rice', 'oryza': 'rice',
              'corn': 'corn', 'maize': 'corn', '玉米': 'corn',
              'wheat': 'wheat', '小麦': 'wheat', 'soybean': 'soybean', '大豆': 'soybean',
              'barley': 'barley', '大麦': 'barley', 'sorghum': 'sorghum', '高粱': 'sorghum'}


def decimal_text(value):
    text = format(value, 'f')
    return (text.rstrip('0').rstrip('.') if '.' in text else text) or '0'


def _cell(value):
    if value is None or value == '':
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    return str(value)


def _xlsx_numbers(raw, sheet):
    """Read stored numeric tokens without routing them through binary floats."""
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
    main = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    rel = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        selected = next(x for x in workbook.find(main + 'sheets') if x.get('name') == sheet)
        relationships = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        target = next(x for x in relationships if x.get('Id') == selected.get(rel + 'id'))
        if target.get('TargetMode') == 'External':
            raise ValueError('不支持外部 XLSX 工作表')
        location = target.get('Target', '')
        location = posixpath.normpath(location.lstrip('/') if location.startswith('/') else 'xl/' + location)
        if not location.startswith('xl/'):
            raise ValueError('XLSX 工作表路径无效')
        worksheet = ET.fromstring(archive.read(location))
        numbers = {}
        row_index = 0
        for row in worksheet.find(main + 'sheetData'):
            row_index = int(row.get('r', row_index + 1))
            column_index = 0
            for cell in row:
                coordinate = cell.get('r')
                if coordinate:
                    column, _ = coordinate_from_string(coordinate)
                    column_index = column_index_from_string(column)
                else:
                    column_index += 1
                    coordinate = get_column_letter(column_index) + str(row_index)
                value = cell.find(main + 'v')
                if cell.get('t', 'n') == 'n' and value is not None:
                    numbers[coordinate] = value.text
        return numbers


def read_table(path, encoding=None, sheet=None):
    """Read a single table. Empty cells become None; nonempty text stays intact."""
    path = Path(path)
    raw = path.read_bytes()
    source = {'filename': path.name, 'sha256': hashlib.sha256(raw).hexdigest()}
    if path.suffix.lower() == '.xlsx':
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet is None and len(book.sheetnames) != 1:
                raise ValueError('XLSX 含多张工作表，请用 --sheet 明确选择')
            selected = sheet or book.sheetnames[0]
            if selected not in book.sheetnames:
                raise ValueError('指定的 XLSX 工作表不存在')
            numeric_tokens = _xlsx_numbers(raw, selected)
            records = []
            for cells in book[selected].iter_rows():
                row = []
                for cell in cells:
                    if cell.data_type == 'f':
                        raise ValueError('XLSX 含公式，请先导出已计算的纯值表格')
                    value = _cell(cell.value)
                    if cell.data_type == 'n' and cell.value is not None and not cell.is_date:
                        if getattr(cell, 'coordinate', None) not in numeric_tokens:
                            raise ValueError('无法关联 XLSX 原始数值，拒绝使用可能丢精度的浮点替代值')
                        value = numeric_tokens[cell.coordinate]
                        if value is not None and re.fullmatch(r'0{2,}', cell.number_format):
                            number = Decimal(value)
                            if number == number.to_integral_value():
                                value = str(int(number)).zfill(len(cell.number_format))
                    row.append(value)
                records.append(row)
            source.update(encoding='xlsx', sheet=selected)
        finally:
            book.close()
    elif path.suffix.lower() in {'.csv', '.tsv', '.txt'}:
        selected_encoding = encoding or ('utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig')
        try:
            content = raw.decode(selected_encoding)
        except (UnicodeDecodeError, LookupError) as exc:
            raise ValueError('无法可靠解码文件，请明确 --encoding，例如 gb18030') from exc
        if '\x00' in content:
            raise ValueError('文件包含 NUL，可能编码错误')
        delimiter = '\t' if path.suffix.lower() == '.tsv' else ','
        if path.suffix.lower() != '.tsv':
            try:
                delimiter = csv.Sniffer().sniff(content[:65536], delimiters=',\t;').delimiter
            except csv.Error:
                pass  # Single-column CSV and malformed rows still receive strict checks.
        try:
            records = [[_cell(v) for v in row] for row in csv.reader(io.StringIO(content, newline=''), delimiter=delimiter, strict=True)]
        except csv.Error as exc:
            raise ValueError('CSV 引号或记录格式错误') from exc
        source.update(encoding=selected_encoding, delimiter=delimiter)
    else:
        raise ValueError('支持 CSV、TSV、TXT 和单工作表 XLSX')
    if not records or not records[0]:
        raise ValueError('数据文件为空或没有表头')
    headers = records[0]
    if any(h is None or not h.strip() for h in headers):
        raise ValueError('存在空表头，请明确列名')
    if len({h.strip().casefold() for h in headers}) != len(headers):
        raise ValueError('存在重复表头 duplicate headers')
    rows = []
    for index, row in enumerate(records[1:], start=2):
        if not row:
            raise ValueError(f'第 {index} 条记录为空行，不能静默丢弃')
        if len(row) != len(headers):
            raise ValueError(f'第 {index} 条记录列数不符：{len(row)}，应为 {len(headers)}')
        rows.append(row)
    if not rows:
        raise ValueError('只有表头，没有数据记录')
    return {'headers': headers, 'rows': rows, 'source': source}


def _crop(table, requested):
    observed = set()
    for index, header in enumerate(table['headers']):
        if header.strip().casefold() in {'作物', '作物名称', 'crop', 'crop_name', 'crop name'}:
            raw_crops = {row[index].strip().casefold() for row in table['rows'] if row[index] is not None}
            normalized = {CROP_NAMES.get(x, x) for x in raw_crops}
            if len(normalized) > 1:
                raise ValueError('检测到多作物 mixed crops，请先拆分数据')
            observed |= normalized
    if requested:
        selected = requested.lower()
        if not re.fullmatch(r'[a-z][a-z0-9_]{0,31}', selected):
            raise ValueError('作物前缀须为小写英文字母开头，最多32字符')
        if observed and selected not in observed:
            raise ValueError('指定作物前缀与文件作物字段冲突')
        return selected
    if len(observed) == 1 and next(iter(observed)) in set(CROP_NAMES.values()):
        return next(iter(observed))
    stem = Path(table['source']['filename']).stem.casefold()
    candidates = {normalized for name, normalized in CROP_NAMES.items()
                  if (name in stem if not name.isascii() else re.search(r'(?:^|[^a-z])' + re.escape(name) + r'(?:$|[^a-z])', stem))}
    if len(candidates) == 1:
        return candidates.pop()
    raise ValueError('无法可靠识别作物英文前缀，请指定 --crop')


def _name(header, used):
    candidate = header.strip()
    if len(candidate) > 64 or re.search(r'[\x00-\x1f\x7f`]', candidate):
        candidate = 'field_' + hashlib.sha256(header.encode()).hexdigest()[:12]
    if candidate.casefold() in RESERVED:
        candidate = 'source_' + candidate
    initial = candidate
    suffix = 2
    while candidate.casefold() in used:
        candidate = initial[:58] + '_' + str(suffix)
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _unit(header):
    known = r'cm|mm|kg|mg|g|m|%|day|days|d|天|个|厘米|毫米'
    match = re.search(r'[（(\[]\s*(' + known + r')\s*[）)\]]\s*$', header, re.I)
    if not match:
        match = re.search(r'_(' + known + r')$', header, re.I)
    return match.group(1) if match else None


def _quantile(numbers, fraction):
    at = Decimal(len(numbers) - 1) * fraction
    low = int(at)
    high = min(low + 1, len(numbers) - 1)
    return numbers[low] + (numbers[high] - numbers[low]) * (at - low)


def _field(header, name, cells):
    observed = [v for v in cells if v is not None]
    counts = Counter(observed)
    notes = []
    stats = {'count': len(cells), 'observed': len(observed), 'missing': len(cells) - len(observed), 'distinct': len(counts)}
    leading_zero = any(re.fullmatch(r'[+-]?0\d+', v) for v in observed)
    identifier = bool(IDENTIFIER.search(header)) or leading_zero
    numerical = bool(observed) and all(NUMBER.fullmatch(v) and len(v) < 200 for v in observed)
    kind = 'identifier' if identifier else 'categorical' if CATEGORY.search(header) else 'numeric' if numerical else 'text'
    if kind == 'numeric' and COUNT.search(header):
        kind = 'count'
    date_formats = Counter()
    for v in observed:
        if re.fullmatch(r'\d{4}-\d{2}-\d{2}(?:T.*)?', v):
            try:
                datetime.fromisoformat(v)
                date_formats['ISO'] += 1
            except ValueError:
                date_formats['invalid_ISO'] += 1
        elif re.fullmatch(r'\d{1,2}/\d{1,2}', v):
            date_formats['month_day_without_year_unvalidated'] += 1
    if not identifier and date_formats:
        kind = 'date'
        stats['date_formats'] = dict(date_formats)
        stats['other_format_count'] = len(observed) - sum(date_formats.values())
        notes.append('日期以原字符串保存；格式不一致或年份缺失时禁止直接按单一日期格式比较。')
    if kind == 'text' and len(counts) <= 100 and all(len(v) <= 128 for v in observed):
        kind = 'categorical'
    max_len = max((len(v) for v in observed), default=1)
    sql_type = f'VARCHAR({max(1, max_len)})' if max_len <= 255 else 'LONGTEXT'
    if numerical and kind in {'numeric', 'count'}:
        with localcontext() as context:
            context.prec = 100
            try:
                numbers = sorted(Decimal(v) for v in observed)
            except InvalidOperation as exc:
                raise ValueError('无法精确解析数值字段') from exc
            if any(not v.is_finite() or abs(v.adjusted()) > 100 for v in numbers):
                raise ValueError('数值指数超出可安全解析范围')
            scale = max(max(0, -v.as_tuple().exponent) for v in numbers)
            digits = max(max(1, v.adjusted() + 1) for v in numbers)
            precision = digits + scale
            if precision > 65 or scale > 30:
                raise ValueError(f'字段 {header} 超出 MySQL DECIMAL 精度限制，不能静默舍入')
            sql_type = f'DECIMAL({precision},{scale})'
            stats.update(min=decimal_text(numbers[0]), max=decimal_text(numbers[-1]),
                         mean=decimal_text(sum(numbers) / len(numbers)),
                         median=decimal_text(_quantile(numbers, Decimal('0.5'))),
                         p25=decimal_text(_quantile(numbers, Decimal('0.25'))),
                         p75=decimal_text(_quantile(numbers, Decimal('0.75'))))
            stats['zero_count'] = sum(v == 0 for v in numbers)
            stats['suspected_missing_codes'] = {decimal_text(v): sum(x == v for x in numbers)
                                                for v in (Decimal('-9'), Decimal('9999')) if v in numbers}
            if stats['zero_count']:
                notes.append('原始零值已保留；是否合理须结合性状定义。')
            if stats['suspected_missing_codes']:
                notes.append('疑似缺失码已原样保留，未自动改成 NULL。')
    elif not identifier and any(NUMBER.fullmatch(v or '') for v in observed) and not numerical:
        notes.append('数值与非数值混存，按字符串保存，不静默丢弃非数值记录。')
    if name != header:
        notes.append('数据库列名由原表头映射，原始表头已记录。')
    if any(v != v.strip() for v in observed):
        notes.append('原始字符串两端空白已保留；匹配时不能忽略这种差异。')
    unit = _unit(header)
    return {'source_name': header, 'name': name, 'sql_type': sql_type, 'kind': kind,
            'unit': unit, 'unit_evidence': 'column_header' if unit else None,
            'stats': stats, 'values': [{'value': v, 'count': n} for v, n in counts.items()], 'notes': notes}


def profile_table(table, crop=None, cache_contract=None):
    selected_crop = _crop(table, crop)
    used = set()
    fields = [_field(header, _name(header, used), [row[index] for row in table['rows']])
              for index, header in enumerate(table['headers'])]
    if len(fields) > 500:
        raise ValueError('第一版最多支持500列，超过时请先明确拆表方案')
    # Large variable-width schemas otherwise exceed MySQL's aggregate row-size limit.
    varchar_bytes = sum(int(re.search(r'\d+', f['sql_type']).group()) * 4 + 2
                        for f in fields if f['sql_type'].startswith('VARCHAR'))
    if varchar_bytes > 48000:
        for field in sorted(fields, key=lambda f: len(f['sql_type']), reverse=True):
            if field['sql_type'].startswith('VARCHAR'):
                width = int(re.search(r'\d+', field['sql_type']).group())
                field['sql_type'] = 'LONGTEXT'
                varchar_bytes -= width * 4 + 2
                if varchar_bytes <= 48000:
                    break
    contract = dict(cache_contract or {'query_task_type': 'VARCHAR(64)', 'time_column': 'create_at',
                                      'session_id_length': 64, 'confirmed': False})
    warnings = ['统计仅描述当前输入数据，不构成农艺最佳区间。']
    if not contract.get('confirmed'):
        warnings.append('缓存接口尚未确认：当前使用用户口述字段生成离线草案，禁止据此直接发布。')
    return {'format_version': 1, 'crop': selected_crop, 'source': table['source'],
            'row_count': len(table['rows']), 'fields': fields, 'warnings': warnings,
            'cache_contract': contract}
